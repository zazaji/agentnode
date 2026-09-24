from __future__ import annotations
import base64, mimetypes, os, re, shutil, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from .config import AppConfig
from .auth import require
from .fuzzy import best_window
from ..models import Principal

class FileManager:
    def __init__(self,cfg:AppConfig): self.cfg=cfg

    def _roots(self):
        return [Path(x).expanduser().resolve(strict=False) for x in self.cfg.files.allowed_roots]

    def _resolve(self,path:str,write=False)->Path:
        p=Path(path).expanduser().resolve(strict=False)
        roots=self._roots()
        if not any(p==root or root in p.parents for root in roots):
            raise PermissionError(f"path outside allowed roots: {p}")
        return p

    def read(self,path:str,p:Principal,offset=0,limit=None):
        """Backward-compatible byte-range reader."""
        require(p,"file.read"); f=self._resolve(path)
        if f.is_dir():
            return {"path":str(f),"kind":"directory","entries":self.list(str(f),p,1)}
        lim=min(limit or self.cfg.files.max_read_bytes,self.cfg.files.max_read_bytes)
        with f.open('rb') as h: h.seek(max(0,offset)); b=h.read(lim)
        return {"path":str(f),"data":b.decode(errors='replace'),"bytes":len(b),"more":f.stat().st_size>offset+len(b),"offset":offset}

    def read_smart(self,path:str,p:Principal,offset:int=0,length:int=1000,
                   sheet:str|None=None,cell_range:str|None=None,include_binary:bool=False):
        """Content-aware reader modeled after Desktop Commander's file handlers.

        Text is line-paged. DOCX/XLSX are extracted using OOXML. PDF support is
        enabled when pypdf is installed. Images return metadata and optional base64.
        """
        require(p,"file.read"); f=self._resolve(path)
        if f.is_dir(): return {"path":str(f),"kind":"directory","entries":self.list(str(f),p,2)}
        if not f.exists(): raise FileNotFoundError(str(f))
        ext=f.suffix.lower(); mime=mimetypes.guess_type(str(f))[0] or "application/octet-stream"
        length=min(max(1,length),5000); offset=max(0,offset)
        if ext==".docx":
            lines=self._read_docx(f); return self._page_lines(f,"docx",mime,lines,offset,length)
        if ext in {".xlsx",".xlsm"}:
            data=self._read_xlsx(f,sheet,cell_range); lines=[]
            for sname,rows in data.items():
                lines.append(f"## Sheet: {sname}")
                lines.extend("\t".join(str(v) for v in row) for row in rows)
            out=self._page_lines(f,"xlsx",mime,lines,offset,length); out["sheets"]=list(data); return out
        if ext==".pdf":
            try:
                from pypdf import PdfReader
                r=PdfReader(str(f)); pages=[(pg.extract_text() or "") for pg in r.pages]
                sel=pages[offset:offset+length]
                return {"path":str(f),"kind":"pdf","mime":mime,"pages":sel,"offset":offset,
                        "returned":len(sel),"total":len(pages),"more":offset+len(sel)<len(pages),
                        "metadata":dict(r.metadata or {})}
            except ImportError:
                return {"path":str(f),"kind":"pdf","mime":mime,"error":"PDF text extraction requires optional dependency pypdf"}
        if mime.startswith("image/"):
            meta={"path":str(f),"kind":"image","mime":mime,"bytes":f.stat().st_size}
            try:
                from PIL import Image
                with Image.open(f) as im: meta.update({"width":im.width,"height":im.height,"mode":im.mode,"format":im.format})
            except Exception: pass
            if include_binary:
                if f.stat().st_size>self.cfg.files.max_read_bytes: raise ValueError("image exceeds max_read_bytes")
                meta["base64"]=base64.b64encode(f.read_bytes()).decode()
            return meta
        # Detect binary before pretending it is text.
        with f.open("rb") as _h: raw=_h.read(8192)
        binary=(b"\x00" in raw)
        if binary:
            out={"path":str(f),"kind":"binary","mime":mime,"bytes":f.stat().st_size}
            if include_binary and f.stat().st_size<=self.cfg.files.max_read_bytes:
                out["base64"]=base64.b64encode(f.read_bytes()).decode()
            return out
        # Stream text paging so huge logs do not have to be loaded into memory.
        chunk=[]; total=0; more=False
        with f.open("r",encoding="utf-8",errors="replace") as h:
            for idx,line in enumerate(h):
                total=idx+1
                if idx < offset: continue
                if len(chunk) < length: chunk.append(line.rstrip("\r\n"))
                else:
                    more=True
                    # Continue counting only for reasonably sized files; exact total
                    # is not worth scanning multi-GB logs after the requested page.
                    if f.stat().st_size > 32_000_000: break
            if not more and total > offset+len(chunk): more=True
        return {"path":str(f),"kind":"text","mime":mime,"content":"\n".join(chunk),"lines":chunk,
                "offset":offset,"returned":len(chunk),"total":None if (more and f.stat().st_size>32_000_000) else total,"more":more}

    @staticmethod
    def _page_lines(f:Path,kind:str,mime:str,lines:list[str],offset:int,length:int):
        chunk=lines[offset:offset+length]
        return {"path":str(f),"kind":kind,"mime":mime,"content":"\n".join(chunk),"lines":chunk,
                "offset":offset,"returned":len(chunk),"total":len(lines),"more":offset+len(chunk)<len(lines)}

    @staticmethod
    def _read_docx(f:Path)->list[str]:
        with zipfile.ZipFile(f) as z:
            root=ET.fromstring(z.read("word/document.xml")); ns="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            out=[]
            for p in root.iter(ns+"p"):
                text="".join((t.text or "") for t in p.iter(ns+"t"))
                if text: out.append(text)
            return out

    @staticmethod
    def _read_xlsx(f:Path,sheet:str|None=None,cell_range:str|None=None)->dict[str,list[list[str]]]:
        # Prefer openpyxl for correct formulas/dates/ranges; provide OOXML fallback.
        try:
            from openpyxl import load_workbook
            wb=load_workbook(f,read_only=True,data_only=False)
            try:
                names=[sheet] if sheet else wb.sheetnames
                out={}
                for name in names:
                    ws=wb[name]
                    cells=ws[cell_range] if cell_range else ws.iter_rows()
                    if cell_range and not isinstance(cells,tuple): cells=((cells,),)
                    rows=[]
                    for row in cells:
                        if not isinstance(row,(tuple,list)): row=(row,)
                        rows.append(["" if c.value is None else str(c.value) for c in row])
                        if len(rows)>=10000: break
                    out[name]=rows
                return out
            finally:
                wb.close()
        except Exception:
            # Invalid/minimal OOXML or optional dependency unavailable: fall back
            # to direct OOXML extraction instead of failing the whole read.
            pass
        with zipfile.ZipFile(f) as z:
            strings=[]
            if "xl/sharedStrings.xml" in z.namelist():
                r=ET.fromstring(z.read("xl/sharedStrings.xml"))
                strings=["".join(t.text or "" for t in si.iter() if t.tag.endswith("}t")) for si in r]
            out={}
            for idx,name in enumerate(sorted(n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")),1):
                sname=f"Sheet{idx}"; rows=[]
                r=ET.fromstring(z.read(name))
                for row in (x for x in r.iter() if x.tag.endswith("}row")):
                    vals=[]
                    for c in (x for x in row if x.tag.endswith("}c")):
                        typ=c.attrib.get("t"); v=next((x.text for x in c if x.tag.endswith("}v")),None)
                        if v is None: v="".join(x.text or "" for x in c.iter() if x.tag.endswith("}t"))
                        if v and typ=="s":
                            try:v=strings[int(v)]
                            except Exception:pass
                        vals.append(v or "")
                    rows.append(vals)
                out[sname]=rows
            return out

    def read_many(self,paths:list[str],p:Principal,smart:bool=True,limit_per_file:int=1000):
        require(p,"file.read"); out=[]
        for path in paths[:50]:
            try: out.append(self.read_smart(path,p,length=limit_per_file) if smart else self.read(path,p))
            except Exception as e: out.append({"path":path,"error":f"{type(e).__name__}: {e}"})
        return out

    def write(self,path:str,data:str,p:Principal,atomic=True,mode:str="rewrite"):
        require(p,"file.write"); f=self._resolve(path,True); b=data.encode()
        if len(b)>self.cfg.files.max_write_bytes: raise ValueError("write too large")
        f.parent.mkdir(parents=True,exist_ok=True)
        if mode=="append":
            with f.open("ab") as h:h.write(b)
        elif mode=="rewrite":
            if atomic:
                tmp=f.with_name(f.name+".agentnode.tmp"); tmp.write_bytes(b); os.replace(tmp,f)
            else:f.write_bytes(b)
        else: raise ValueError("mode must be rewrite or append")
        return {"path":str(f),"bytes":len(b),"mode":mode}

    def mkdir(self,path:str,p:Principal):
        require(p,"file.write"); f=self._resolve(path,True); f.mkdir(parents=True,exist_ok=True); return {"path":str(f),"created":True}

    def move(self,source:str,destination:str,p:Principal,overwrite:bool=False):
        require(p,"file.write"); src=self._resolve(source,True); dst=self._resolve(destination,True)
        if dst.exists() and not overwrite: raise FileExistsError(str(dst))
        dst.parent.mkdir(parents=True,exist_ok=True)
        if overwrite and dst.exists():
            if dst.is_dir(): shutil.rmtree(dst)
            else: dst.unlink()
        return {"source":str(src),"destination":str(shutil.move(str(src),str(dst)))}

    def info(self,path:str,p:Principal):
        require(p,"file.read"); f=self._resolve(path); st=f.stat(); mime=mimetypes.guess_type(str(f))[0]
        out={"path":str(f),"name":f.name,"exists":True,"type":"directory" if f.is_dir() else "file",
             "size":st.st_size,"created":st.st_ctime,"modified":st.st_mtime,"mode":oct(st.st_mode & 0o777),"mime":mime,
             "symlink":f.is_symlink()}
        if f.suffix.lower() in {".xlsx",".xlsm"}:
            try: out["sheets"]=list(self._read_xlsx(f).keys())
            except Exception: pass
        return out

    def list(self,path:str,p:Principal,depth=1):
        require(p,"file.read"); root=self._resolve(path); out=[]
        if not root.exists(): raise FileNotFoundError(str(root))
        for cur,dirs,files in os.walk(root):
            rel=Path(cur).relative_to(root); d=len(rel.parts)
            if d>=depth: dirs[:]=[]
            for n in dirs+files:
                q=Path(cur)/n
                try: size=q.stat().st_size if q.is_file() else None
                except OSError: size=None
                out.append({"path":str(q),"relative":str(q.relative_to(root)),"dir":q.is_dir(),"size":size})
                if len(out)>=2000:return out
        return out

    def replace(self,path:str,old:str,new:str,p:Principal,expected=1,fuzzy:bool=False,min_similarity:float=0.86,dry_run:bool=False):
        require(p,"file.write"); f=self._resolve(path,True)
        if f.stat().st_size > self.cfg.files.max_write_bytes: raise ValueError("file too large for replace operation")
        s=f.read_text('utf-8'); n=s.count(old)
        if n==expected:
            changed=s.replace(old,new)
            if dry_run:return {"path":str(f),"matches":n,"exact":True,"would_change":True}
            r=self.write(str(f),changed,p); r.update({"matches":n,"exact":True}); return r
        suggestion=best_window(s,old)
        if not fuzzy or suggestion["similarity"]<min_similarity:
            raise ValueError(f"expected {expected} exact matches, found {n}; closest similarity={suggestion['similarity']}, closest={suggestion['value'][:300]!r}")
        if expected!=1: raise ValueError("fuzzy replacement supports expected=1 only")
        start,end=suggestion["start"],suggestion["end"]; changed=s[:start]+new+s[end:]
        if dry_run:return {"path":str(f),"exact":False,"similarity":suggestion["similarity"],"matched":suggestion["value"],"would_change":True}
        r=self.write(str(f),changed,p); r.update({"exact":False,"similarity":suggestion["similarity"],"matched":suggestion["value"]}); return r

    def search(self,root:str,pattern:str,p:Principal,limit=200):
        """Compatibility one-shot search. Prefer SearchManager for large searches."""
        require(p,"file.read"); base=self._resolve(root); out=[]; rg=shutil.which('rg')
        if rg:
            import subprocess,json
            cp=subprocess.run([rg,'--json',pattern,str(base)],capture_output=True,text=True,timeout=30)
            for line in cp.stdout.splitlines():
                try:
                    j=json.loads(line); d=j.get('data',{})
                    if j.get('type')=='match': out.append({"path":d['path']['text'],"line":d['line_number'],"text":d['lines']['text'].rstrip()})
                except Exception: pass
                if len(out)>=limit:break
            return out
        rx=re.compile(pattern)
        for f in base.rglob('*'):
            if not f.is_file(): continue
            try:
                for i,line in enumerate(f.read_text('utf-8',errors='ignore').splitlines(),1):
                    if rx.search(line): out.append({"path":str(f),"line":i,"text":line[:1000]})
                    if len(out)>=limit:return out
            except OSError: pass
        return out
