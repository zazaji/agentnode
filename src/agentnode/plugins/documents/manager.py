from __future__ import annotations
from pathlib import Path
from typing import Any
import shutil
from ...core.auth import require


class DocumentManager:
    """Optional rich document operations. Dependencies are loaded lazily."""
    def __init__(self, files):
        self.files = files

    def _path(self, path: str, p, write: bool = False) -> Path:
        # Reuse FileManager's path guard rather than creating a second sandbox.
        return self.files._resolve(path, write=write)

    def capabilities(self) -> dict[str, Any]:
        out = {}
        for mod in ("openpyxl", "docx", "pypdf"):
            try:
                __import__(mod); out[mod] = True
            except Exception:
                out[mod] = False
        return out

    def xlsx_read(self, path: str, p, sheet: str | None = None, cell_range: str | None = None, formulas: bool = True):
        require(p, "document.read")
        fp = self._path(path, p)
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError("openpyxl is required; install agentnode[documents]") from e
        wb = openpyxl.load_workbook(fp, data_only=not formulas, read_only=True)
        try:
            ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
            cells = ws[cell_range] if cell_range else ws.iter_rows()
            rows = []
            for row in cells:
                if not isinstance(row, (list, tuple)): row = [row]
                rows.append([c.value for c in row])
                if len(rows) >= 5000: break
            sheets = list(wb.sheetnames)
            title = ws.title
        finally:
            # read_only workbooks keep the zip stream open; on Windows an open
            # handle blocks later atomic replace of the same file.
            wb.close()
        return {"kind":"xlsx","path":str(fp),"sheet":title,"sheets":sheets,"rows":rows,"truncated":len(rows)>=5000}

    def xlsx_set_range(self, path: str, p, sheet: str, start_cell: str, values: list[list[Any]], create: bool = False):
        require(p, "document.write")
        fp = self._path(path, p, write=True)
        try:
            import openpyxl
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        except ImportError as e:
            raise RuntimeError("openpyxl is required; install agentnode[documents]") from e
        if fp.exists(): wb = openpyxl.load_workbook(fp)
        elif create: wb = openpyxl.Workbook()
        else: raise FileNotFoundError(fp)
        try:
            ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
            col_s, row0 = coordinate_from_string(start_cell); col0 = column_index_from_string(col_s)
            for r_off, row in enumerate(values):
                for c_off, value in enumerate(row): ws.cell(row=row0+r_off, column=col0+c_off, value=value)
            tmp = fp.with_suffix(fp.suffix + ".tmp")
            wb.save(tmp)
        finally:
            wb.close()
        tmp.replace(fp)
        return {"ok":True,"path":str(fp),"sheet":sheet,"start_cell":start_cell,"rows":len(values),"cols":max((len(x) for x in values), default=0)}

    def xlsx_append_rows(self, path: str, p, sheet: str, values: list[list[Any]]):
        require(p, "document.write")
        fp = self._path(path, p, write=True)
        try: import openpyxl
        except ImportError as e: raise RuntimeError("openpyxl is required; install agentnode[documents]") from e
        wb = openpyxl.load_workbook(fp)
        try:
            ws = wb[sheet]
            for row in values: ws.append(row)
            tmp = fp.with_suffix(fp.suffix + ".tmp"); wb.save(tmp)
        finally:
            wb.close()
        tmp.replace(fp)
        return {"ok":True,"path":str(fp),"sheet":sheet,"appended":len(values)}

    def docx_read(self, path: str, p):
        require(p, "document.read")
        fp=self._path(path,p)
        try:
            from docx import Document
        except ImportError as e: raise RuntimeError("python-docx is required; install agentnode[documents]") from e
        d=Document(fp)
        return {"kind":"docx","path":str(fp),"paragraphs":[x.text for x in d.paragraphs],"tables":[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]}

    def docx_replace(self, path: str, p, old: str, new: str, expected: int | None = None):
        require(p, "document.write")
        fp=self._path(path,p,write=True)
        try:
            from docx import Document
        except ImportError as e: raise RuntimeError("python-docx is required; install agentnode[documents]") from e
        d=Document(fp); count=0
        for para in d.paragraphs:
            if old in para.text:
                # Preserve paragraph-level style where possible; run-level formatting may change only in edited paragraph.
                for run in para.runs:
                    n=run.text.count(old)
                    if n: run.text=run.text.replace(old,new); count+=n
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            n=run.text.count(old)
                            if n: run.text=run.text.replace(old,new); count+=n
        if expected is not None and count != expected: raise ValueError(f"expected {expected} replacements, found {count}")
        tmp=fp.with_suffix(fp.suffix+'.tmp'); d.save(tmp); tmp.replace(fp)
        return {"ok":True,"path":str(fp),"replacements":count}

    def docx_create(self, path: str, p, paragraphs: list[str], title: str | None = None):
        require(p, "document.write")
        fp=self._path(path,p,write=True); fp.parent.mkdir(parents=True,exist_ok=True)
        try:
            from docx import Document
        except ImportError as e: raise RuntimeError("python-docx is required; install agentnode[documents]") from e
        d=Document()
        if title: d.add_heading(title,0)
        for item in paragraphs: d.add_paragraph(item)
        tmp=fp.with_suffix(fp.suffix+'.tmp'); d.save(tmp); tmp.replace(fp)
        return {"ok":True,"path":str(fp),"paragraphs":len(paragraphs)}

    def pdf_read(self, path: str, p, start_page: int = 0, pages: int = 20):
        require(p, "document.read")
        fp=self._path(path,p)
        try: from pypdf import PdfReader
        except ImportError as e: raise RuntimeError("pypdf is required; install agentnode[documents]") from e
        r=PdfReader(fp); end=min(len(r.pages), start_page+pages)
        return {"kind":"pdf","path":str(fp),"page_count":len(r.pages),"start_page":start_page,"pages":[{"page":i,"text":r.pages[i].extract_text() or ""} for i in range(start_page,end)]}

    def pdf_merge(self, inputs: list[str], output: str, p):
        require(p, "document.write")
        try: from pypdf import PdfReader, PdfWriter
        except ImportError as e: raise RuntimeError("pypdf is required; install agentnode[documents]") from e
        out=self._path(output,p,write=True); writer=PdfWriter()
        for item in inputs:
            src=self._path(item,p)
            for page in PdfReader(src).pages: writer.add_page(page)
        tmp=out.with_suffix(out.suffix+'.tmp')
        with tmp.open('wb') as f: writer.write(f)
        tmp.replace(out)
        return {"ok":True,"output":str(out),"inputs":len(inputs),"pages":len(writer.pages)}
